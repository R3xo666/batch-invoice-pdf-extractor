"""Excel workbook generation engine using openpyxl with professional styling and audit reporting."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.models import (
    AuditRecord,
    BatchSummary,
    FLAGGED_FOR_REVIEW,
    FLAGGED_SCANNED,
    FLAGGED_UNREADABLE,
    InvoiceRecord,
    ProcessingStatus,
)

logger = logging.getLogger(__name__)

# Color Palette Constants
COLOR_NAVY_HEADER = "1B365D"
COLOR_WHITE = "FFFFFF"
COLOR_ZEBRA_LIGHT = "F8F9FA"
COLOR_FLAGGED_YELLOW = "FFF3CD"
COLOR_FLAGGED_TEXT = "856404"
COLOR_SUCCESS_GREEN = "D4EDDA"
COLOR_SUCCESS_TEXT = "155724"
COLOR_FAILED_RED = "F8D7DA"
COLOR_FAILED_TEXT = "721C24"
COLOR_BORDER = "D0D5DD"
COLOR_KPI_BG = "E9ECEF"

SUMMARY_COLUMNS = [
    "Filename",
    "Vendor Name",
    "Invoice Number",
    "Invoice Date",
    "Subtotal",
    "Tax / VAT",
    "Total Due",
    "Status",
]

AUDIT_COLUMNS = [
    "Timestamp",
    "Filename",
    "Pages",
    "Processing Status",
    "Notes / Diagnostics",
]


class ExcelExporter:
    """Creates a beautifully styled, multi-sheet Excel workbook from invoice extraction results."""

    def __init__(self) -> None:
        """Initialize style objects for cell rendering."""
        self.thin_border = Border(
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
            top=Side(style="thin", color=COLOR_BORDER),
            bottom=Side(style="thin", color=COLOR_BORDER),
        )
        self.header_font = Font(name="Calibri", size=11, bold=True, color=COLOR_WHITE)
        self.header_fill = PatternFill(
            start_color=COLOR_NAVY_HEADER, end_color=COLOR_NAVY_HEADER, fill_type="solid"
        )
        self.regular_font = Font(name="Calibri", size=10)
        self.bold_font = Font(name="Calibri", size=10, bold=True)
        self.zebra_fill = PatternFill(
            start_color=COLOR_ZEBRA_LIGHT, end_color=COLOR_ZEBRA_LIGHT, fill_type="solid"
        )
        self.white_fill = PatternFill(
            start_color=COLOR_WHITE, end_color=COLOR_WHITE, fill_type="solid"
        )
        self.flagged_fill = PatternFill(
            start_color=COLOR_FLAGGED_YELLOW, end_color=COLOR_FLAGGED_YELLOW, fill_type="solid"
        )
        self.flagged_font = Font(name="Calibri", size=10, color=COLOR_FLAGGED_TEXT, bold=True)

        self.success_fill = PatternFill(
            start_color=COLOR_SUCCESS_GREEN, end_color=COLOR_SUCCESS_GREEN, fill_type="solid"
        )
        self.success_font = Font(name="Calibri", size=10, color=COLOR_SUCCESS_TEXT, bold=True)

        self.failed_fill = PatternFill(
            start_color=COLOR_FAILED_RED, end_color=COLOR_FAILED_RED, fill_type="solid"
        )
        self.failed_font = Font(name="Calibri", size=10, color=COLOR_FAILED_TEXT, bold=True)

    def _auto_fit_columns(self, sheet: Any, max_width_limit: int = 50) -> None:
        """Calculate and set dynamic column widths based on contents."""
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_width_limit)

    def _format_summary_sheet(
        self, sheet: Any, records: List[InvoiceRecord]
    ) -> None:
        """Populate and format the main Invoice Summary sheet."""
        sheet.title = "Invoice Summary"
        sheet.views.sheetView[0].showGridLines = True
        sheet.freeze_panes = "A2"

        # Write Header
        for col_idx, col_name in enumerate(SUMMARY_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=col_name)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.thin_border
        sheet.row_dimensions[1].height = 26

        # Write Data Rows
        flagged_markers = (FLAGGED_FOR_REVIEW, FLAGGED_SCANNED, FLAGGED_UNREADABLE)

        for row_idx, record in enumerate(records, start=2):
            row_data = record.to_summary_row()
            is_even = (row_idx % 2 == 0)
            default_fill = self.white_fill if is_even else self.zebra_fill

            for col_idx, raw_val in enumerate(row_data, start=1):
                # Neutralize potential formula injection (CWE-1236) on string values
                val = raw_val
                if isinstance(val, str) and val and val[0] in ("=", "+", "-", "@", "\t", "\r"):
                    val = f"'{val}"

                cell = sheet.cell(row=row_idx, column=col_idx, value=val)
                cell.font = self.regular_font
                cell.fill = default_fill
                cell.border = self.thin_border
                cell.alignment = Alignment(vertical="center")

                # Column-specific formatting
                # Date (Column 4)
                if col_idx == 4:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Financial columns (Subtotal: 5, Tax: 6, Total Due: 7)
                if col_idx in (5, 6, 7):
                    if isinstance(val, (int, float)):
                        cell.number_format = "$#,##0.00"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                # Status column (Column 8)
                if col_idx == 8:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if val == ProcessingStatus.SUCCESS.value:
                        cell.fill = self.success_fill
                        cell.font = self.success_font
                    elif val == ProcessingStatus.FLAGGED.value:
                        cell.fill = self.flagged_fill
                        cell.font = self.flagged_font
                    elif val == ProcessingStatus.FAILED.value:
                        cell.fill = self.failed_fill
                        cell.font = self.failed_font

                # Highlight flagged / unlocatable values
                if isinstance(val, str) and any(marker in val for marker in flagged_markers):
                    cell.fill = self.flagged_fill
                    cell.font = self.flagged_font

            sheet.row_dimensions[row_idx].height = 20

        self._auto_fit_columns(sheet)

    def _format_audit_sheet(
        self, sheet: Any, audit_records: List[AuditRecord], summary: BatchSummary
    ) -> None:
        """Populate and format the Audit & Errors sheet with top KPI cards."""
        sheet.title = "Audit & Errors"
        sheet.views.sheetView[0].showGridLines = True

        # Render KPI Cards in Rows 1-3
        kpi_metrics = [
            ("Total Files Processed", summary.total_files, "B2", "B3"),
            ("Successfully Extracted", summary.total_successful, "C2", "C3"),
            ("Flagged for Review", summary.total_flagged, "D2", "D3"),
            ("Extraction Failures", summary.total_failed, "E2", "E3"),
        ]

        # Title Row
        title_cell = sheet.cell(row=1, column=2, value="Batch Ingestion Audit Dashboard")
        title_cell.font = Font(name="Calibri", size=14, bold=True, color=COLOR_NAVY_HEADER)

        for label, val, top_ref, val_ref in kpi_metrics:
            t_cell = sheet[top_ref]
            t_cell.value = label
            t_cell.font = Font(name="Calibri", size=9, bold=True, color="555555")
            t_cell.alignment = Alignment(horizontal="center", vertical="center")
            t_cell.fill = PatternFill(start_color=COLOR_KPI_BG, end_color=COLOR_KPI_BG, fill_type="solid")
            t_cell.border = self.thin_border

            v_cell = sheet[val_ref]
            v_cell.value = val
            v_cell.font = Font(name="Calibri", size=16, bold=True, color=COLOR_NAVY_HEADER)
            v_cell.alignment = Alignment(horizontal="center", vertical="center")
            v_cell.border = self.thin_border

        sheet.row_dimensions[2].height = 18
        sheet.row_dimensions[3].height = 28

        # Audit Table Header starts at Row 5
        start_row = 5
        for col_idx, col_name in enumerate(AUDIT_COLUMNS, start=1):
            cell = sheet.cell(row=start_row, column=col_idx, value=col_name)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border
        sheet.row_dimensions[start_row].height = 24

        # Audit Records Data
        for idx, audit in enumerate(audit_records, start=start_row + 1):
            row_data = audit.to_audit_row()
            is_even = (idx % 2 == 0)
            default_fill = self.white_fill if is_even else self.zebra_fill

            for col_idx, raw_val in enumerate(row_data, start=1):
                # Neutralize potential formula injection (CWE-1236)
                val = raw_val
                if isinstance(val, str) and val and val[0] in ("=", "+", "-", "@", "\t", "\r"):
                    val = f"'{val}"

                cell = sheet.cell(row=idx, column=col_idx, value=val)
                cell.font = self.regular_font
                cell.fill = default_fill
                cell.border = self.thin_border
                cell.alignment = Alignment(vertical="center")

                # Center align Timestamp, Pages, Status
                if col_idx in (1, 3, 4):
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Color Status
                if col_idx == 4:
                    if val == ProcessingStatus.SUCCESS.value:
                        cell.fill = self.success_fill
                        cell.font = self.success_font
                    elif val == ProcessingStatus.FLAGGED.value:
                        cell.fill = self.flagged_fill
                        cell.font = self.flagged_font
                    elif val == ProcessingStatus.FAILED.value:
                        cell.fill = self.failed_fill
                        cell.font = self.failed_font

            sheet.row_dimensions[idx].height = 20

        self._auto_fit_columns(sheet, max_width_limit=65)

    def export(
        self,
        output_path: Path,
        records: List[InvoiceRecord],
        audit_records: List[AuditRecord],
        summary: BatchSummary,
    ) -> Path:
        """Create and save the complete Excel workbook.

        Args:
            output_path: Destination path for .xlsx file.
            records: Processed invoice records.
            audit_records: Per-file audit logs.
            summary: Aggregate batch KPIs.

        Returns:
            Resolved output file Path.
        """
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()

        # Sheet 1: Invoice Summary
        summary_sheet = workbook.active
        self._format_summary_sheet(summary_sheet, records)

        # Sheet 2: Audit & Errors
        audit_sheet = workbook.create_sheet(title="Audit & Errors")
        self._format_audit_sheet(audit_sheet, audit_records, summary)

        workbook.save(output_path)
        logger.info("Saved formatted Excel workbook to: %s", output_path)
        return output_path
