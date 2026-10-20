"""Unit tests for the Invoice Extractor, Parser, Models, and Excel Exporter."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.exporter import ExcelExporter
from src.extractor import PDFExtractionResult, PDFExtractor
from src.models import (
    AuditRecord,
    BatchSummary,
    FLAGGED_FOR_REVIEW,
    FLAGGED_SCANNED,
    FLAGGED_UNREADABLE,
    InvoiceRecord,
    ProcessingStatus,
)
from src.parser import InvoiceParser


class TestInvoiceModels:
    """Test data model behavior and status evaluation."""

    def test_invoice_record_success_status(self) -> None:
        rec = InvoiceRecord(
            filename="inv1.pdf",
            vendor_name="Acme Corp",
            invoice_number="INV-101",
            invoice_date="2026-10-15",
            subtotal=100.0,
            tax=10.0,
            total_due=110.0,
        )
        assert rec.evaluate_status() == ProcessingStatus.SUCCESS
        assert not rec.is_flagged()

    def test_invoice_record_flagged_status(self) -> None:
        rec = InvoiceRecord(
            filename="inv2.pdf",
            vendor_name=FLAGGED_FOR_REVIEW,
            invoice_number="INV-102",
            invoice_date="2026-10-15",
            subtotal=100.0,
            tax=0.0,
            total_due=100.0,
        )
        assert rec.is_flagged()
        assert rec.evaluate_status() == ProcessingStatus.FLAGGED

    def test_batch_summary_counters(self) -> None:
        summary = BatchSummary()
        summary.add_record(
            InvoiceRecord(filename="f1.pdf", status=ProcessingStatus.SUCCESS, total_due=100.0)
        )
        summary.add_record(
            InvoiceRecord(filename="f2.pdf", status=ProcessingStatus.FLAGGED, total_due=250.50)
        )
        summary.add_record(
            InvoiceRecord(filename="f3.pdf", status=ProcessingStatus.FAILED, total_due=0.0)
        )

        assert summary.total_files == 3
        assert summary.total_successful == 1
        assert summary.total_flagged == 1
        assert summary.total_failed == 1
        assert summary.total_revenue_due == 350.50


class TestInvoiceParser:
    """Test text parsing, currency cleaning, and regex heuristics."""

    @pytest.fixture
    def parser(self) -> InvoiceParser:
        return InvoiceParser()

    def test_clean_currency_us_format(self, parser: InvoiceParser) -> None:
        assert parser.clean_currency("$1,250.00") == 1250.00
        assert parser.clean_currency(" $ 49.99 ") == 49.99
        assert parser.clean_currency("10,000,500.50") == 10000500.50

    def test_clean_currency_european_format(self, parser: InvoiceParser) -> None:
        assert parser.clean_currency("1.250,00 €") == 1250.00
        assert parser.clean_currency("€ 450,50") == 450.50
        assert parser.clean_currency("208,33 EUR") == 208.33

    def test_clean_currency_invalid(self, parser: InvoiceParser) -> None:
        assert parser.clean_currency("N/A") == FLAGGED_FOR_REVIEW
        assert parser.clean_currency("") == FLAGGED_FOR_REVIEW
        assert parser.clean_currency(None) == FLAGGED_FOR_REVIEW

    def test_parse_date_iso(self, parser: InvoiceParser) -> None:
        assert parser.parse_date("2026-10-15") == "2026-10-15"

    def test_parse_date_slash(self, parser: InvoiceParser) -> None:
        # Standard dayfirst / monthfirst parsing
        date_str = parser.parse_date("18/10/2026")
        assert date_str == "2026-10-18"

    def test_parse_date_alphanumeric(self, parser: InvoiceParser) -> None:
        assert parser.parse_date("24-Oct-2026") == "2026-10-24"
        assert parser.parse_date("October 24, 2026") == "2026-10-24"

    def test_extract_invoice_number(self, parser: InvoiceParser) -> None:
        text = "Invoice Number: INV-2026-8941\nDate: 2026-10-15"
        assert parser.extract_invoice_number(text) == "INV-2026-8941"

        text2 = "Bill No: VDL-2026/OCT/092\nAmount: $100"
        assert parser.extract_invoice_number(text2) == "VDL-2026/OCT/092"

    def test_extract_vendor_name(self, parser: InvoiceParser) -> None:
        text = "Vendor: Nordic Data Systems AB\nInvoice # 123"
        lines = ["Vendor: Nordic Data Systems AB", "Invoice # 123"]
        assert parser.extract_vendor_name(text, lines) == "Nordic Data Systems AB"

        text_header = "Apex Cloud Solutions Inc.\n100 Innovation Way\nInvoice Date: 2026-10-15"
        lines_header = ["Apex Cloud Solutions Inc.", "100 Innovation Way", "Invoice Date: 2026-10-15"]
        assert parser.extract_vendor_name(text_header, lines_header) == "Apex Cloud Solutions Inc."

    def test_parse_unreadable_result(self, parser: InvoiceParser) -> None:
        result = PDFExtractionResult(
            file_path=Path("corrupt.pdf"),
            error="PDF syntax error: incomplete xref",
        )
        record = parser.parse(result)
        assert record.status == ProcessingStatus.FAILED
        assert record.vendor_name == FLAGGED_UNREADABLE
        assert record.invoice_number == FLAGGED_UNREADABLE

    def test_parse_scanned_result(self, parser: InvoiceParser) -> None:
        result = PDFExtractionResult(
            file_path=Path("scanned.pdf"),
            is_scanned=True,
            total_pages=1,
            full_text="",
        )
        record = parser.parse(result)
        assert record.status == ProcessingStatus.FLAGGED
        assert record.vendor_name == FLAGGED_SCANNED


class TestExcelExporter:
    """Test openpyxl workbook export and sheet verification."""

    def test_export_creates_valid_workbook(self, tmp_path: Path) -> None:
        output_file = tmp_path / "monthly_summary.xlsx"
        records = [
            InvoiceRecord(
                filename="inv1.pdf",
                vendor_name="Acme Inc",
                invoice_number="INV-001",
                invoice_date="2026-10-01",
                subtotal=100.0,
                tax=10.0,
                total_due=110.0,
                status=ProcessingStatus.SUCCESS,
            ),
            InvoiceRecord(
                filename="inv2.pdf",
                vendor_name=FLAGGED_FOR_REVIEW,
                invoice_number="INV-002",
                invoice_date="2026-10-02",
                subtotal=200.0,
                tax=0.0,
                total_due=200.0,
                status=ProcessingStatus.FLAGGED,
            ),
        ]
        audit_records = [
            AuditRecord(filename="inv1.pdf", pages=1, status=ProcessingStatus.SUCCESS, notes="Clean"),
            AuditRecord(filename="inv2.pdf", pages=1, status=ProcessingStatus.FLAGGED, notes="Missing vendor"),
        ]
        summary = BatchSummary(total_files=2, total_successful=1, total_flagged=1, total_revenue_due=310.0)

        exporter = ExcelExporter()
        exporter.export(output_file, records, audit_records, summary)

        assert output_file.exists()

        # Verify sheets with openpyxl
        wb = load_workbook(output_file)
        assert "Invoice Summary" in wb.sheetnames
        assert "Audit & Errors" in wb.sheetnames

        summary_sheet = wb["Invoice Summary"]
        assert summary_sheet.cell(row=1, column=1).value == "Filename"
        assert summary_sheet.cell(row=2, column=2).value == "Acme Inc"
        assert summary_sheet.cell(row=3, column=2).value == FLAGGED_FOR_REVIEW

        audit_sheet = wb["Audit & Errors"]
        assert audit_sheet["B3"].value == 2  # Total files KPI
        assert audit_sheet["C3"].value == 1  # Successful KPI

    def test_export_sanitizes_formula_injection(self, tmp_path: Path) -> None:
        """Verify potential Excel formula injection (CWE-1236) is neutralized with leading quote."""
        output_file = tmp_path / "formula_test.xlsx"
        records = [
            InvoiceRecord(
                filename="=cmd|' /C calc'!A0.pdf",
                vendor_name="+Acme Inc",
                invoice_number="@INV-001",
                invoice_date="2026-10-01",
                subtotal=100.0,
                tax=0.0,
                total_due=100.0,
                status=ProcessingStatus.SUCCESS,
            ),
        ]
        audit_records = [
            AuditRecord(
                filename="=cmd|' /C calc'!A0.pdf",
                pages=1,
                status=ProcessingStatus.SUCCESS,
                notes="=SUM(A1:A10)",
            )
        ]
        summary = BatchSummary(total_files=1, total_successful=1, total_revenue_due=100.0)

        exporter = ExcelExporter()
        exporter.export(output_file, records, audit_records, summary)

        wb = load_workbook(output_file)
        summary_sheet = wb["Invoice Summary"]
        # Filename, vendor, invoice number must start with single quote
        assert summary_sheet.cell(row=2, column=1).value.startswith("'")
        assert summary_sheet.cell(row=2, column=2).value.startswith("'")
        assert summary_sheet.cell(row=2, column=3).value.startswith("'")

        audit_sheet = wb["Audit & Errors"]
        assert audit_sheet.cell(row=6, column=2).value.startswith("'")
        assert audit_sheet.cell(row=6, column=5).value.startswith("'")


class TestExtractorSecurityLimits:
    """Test resource limits and DoS protections in PDFExtractor."""

    def test_extractor_file_size_limit(self, tmp_path: Path) -> None:
        large_file = tmp_path / "oversized.pdf"
        # Create a sparse file or mock stat exceeding limit
        extractor = PDFExtractor()
        extractor.MAX_FILE_SIZE_BYTES = 100  # Set low for test
        large_file.write_bytes(b"%PDF-1.4" + b"0" * 200)

        result = extractor.extract(large_file)
        assert not result.is_valid
        assert "exceeds maximum allowed size" in (result.error or "")

